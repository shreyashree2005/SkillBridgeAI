"""
Resume analysis (rule-based, no machine-learning model).

1. Extract text from PDF (pypdf), DOCX (python-docx) or TXT.
2. Detect sections by heading keywords.
3. Extract skills with the taxonomy matcher, estimate years of experience
   from date ranges, and pick out education lines.
4. Score quality 0-100 with a transparent rubric (each criterion is returned
   with its points so users can see exactly why they got their score).
"""
from __future__ import annotations

import io
import re
from datetime import date

from ..taxonomy import get_taxonomy

MAX_PDF_PAGES = 15
MAX_CHARS = 100_000


class ResumeError(ValueError):
    """Raised for unreadable or unsupported files (reported to the user as a 400)."""


# ---------------------------------------------------------------- extraction
def extract_text(data: bytes, filename: str) -> str:
    name = filename.lower()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ResumeError("This PDF is password-protected.")
            pages = reader.pages[:MAX_PDF_PAGES]
            text = "\n".join((p.extract_text() or "") for p in pages)
        elif name.endswith(".docx"):
            from docx import Document

            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            text = "\n".join(parts)
        elif name.endswith(".txt"):
            text = data.decode("utf-8", errors="replace")
        else:
            raise ResumeError("Unsupported file type. Use PDF, DOCX or TXT.")
    except ResumeError:
        raise
    except Exception as exc:  # corrupted files raise many different library errors
        raise ResumeError("Could not read this file. It may be corrupted.") from exc

    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_CHARS]


# ---------------------------------------------------------------- parsing
SECTION_KEYWORDS = {
    "summary": ["summary", "profile", "objective", "about me", "professional summary", "career objective"],
    "skills": ["skills", "technical skills", "core competencies", "technologies", "tech stack", "tools"],
    "experience": ["experience", "work experience", "professional experience", "employment", "internships", "internship", "work history"],
    "projects": ["projects", "academic projects", "personal projects", "key projects"],
    "education": ["education", "academic background", "qualifications", "academics"],
    "certifications": ["certifications", "certificates", "courses", "licenses"],
    "achievements": ["achievements", "awards", "honors", "accomplishments"],
    "leadership": ["leadership", "extracurricular", "volunteer", "activities", "positions of responsibility"],
    "publications": ["publications", "research"],
}
_HEADING_LOOKUP = {kw: sec for sec, kws in SECTION_KEYWORDS.items() for kw in kws}

MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
DATE_TOKEN = r"(?:(?P<{p}m>jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s+)?(?P<{p}y>(?:19|20)\d{{2}})"
RANGE_RX = re.compile(
    DATE_TOKEN.format(p="a") + r"\s*(?:-|–|—|to)\s*(?:" + DATE_TOKEN.format(p="b") + r"|(?P<present>present|current|now|ongoing))",
    re.IGNORECASE,
)
YEARS_PHRASE = re.compile(r"(\d{1,2})\+?\s*(?:years|yrs)\s+(?:of\s+)?(?:professional\s+|industry\s+|work\s+)?experience", re.IGNORECASE)
EMAIL_RX = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RX = re.compile(r"(?:\+?\d[\d\s().-]{8,}\d)")
LINK_RX = re.compile(r"(github|linkedin|gitlab|portfolio|behance|dribbble|https?://)", re.IGNORECASE)
DEGREE_RX = re.compile(
    r"\b(b\.?\s?tech|b\.?\s?e\b|b\.?\s?sc|bachelor|m\.?\s?tech|m\.?\s?sc|master|mba|ph\.?\s?d|diploma|associate degree|high school|bca|mca)",
    re.IGNORECASE,
)
ACTION_VERBS = {
    "built", "developed", "designed", "implemented", "led", "created", "improved", "optimized", "optimised",
    "launched", "automated", "reduced", "increased", "deployed", "architected", "analyzed", "analysed",
    "migrated", "refactored", "mentored", "delivered", "integrated", "engineered", "streamlined", "scaled",
}
NUMBER_RX = re.compile(r"\d+(?:\.\d+)?\s*(?:%|x\b|k\b|\+|ms\b|users|students|hours|requests)", re.IGNORECASE)


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw in text.splitlines():
        line = raw.strip()
        key = re.sub(r"[^a-z ]", "", line.lower()).strip()
        if line and len(line.split()) <= 4 and key in _HEADING_LOOKUP:
            current = _HEADING_LOOKUP[key]
            sections.setdefault(current, [])
            continue
        sections.setdefault(current, []).append(raw)
    return {k: "\n".join(v).strip() for k, v in sections.items()}


def _month_index(year: int, month: int) -> int:
    return year * 12 + (month - 1)


def estimate_years(text: str, today: date | None = None) -> float:
    """Merge date ranges (e.g. 'Jan 2022 – Present') into total years of experience."""
    today = today or date.today()
    intervals = []
    for m in RANGE_RX.finditer(text):
        ay = int(m.group("ay"))
        am = MONTHS.get((m.group("am") or "jan")[:3].lower(), 1)
        if m.group("present"):
            by, bm = today.year, today.month
        else:
            by = int(m.group("by"))
            bm = MONTHS.get(m.group("bm")[:3].lower(), 12) if m.group("bm") else 12
        start = _month_index(ay, am)
        end = min(_month_index(by, bm), _month_index(today.year, today.month))  # never count future months
        if 0 <= end - start <= 40 * 12 and by <= today.year + 1:
            intervals.append((start, end + 1))  # inclusive: "2021 - 2022" is 24 months
    intervals.sort()
    total, cur_s, cur_e = 0, None, None
    for s, e in intervals:
        if cur_e is None or s > cur_e:
            if cur_e is not None:
                total += cur_e - cur_s
            cur_s, cur_e = s, e
        else:
            cur_e = max(cur_e, e)
    if cur_e is not None:
        total += cur_e - cur_s
    from_ranges = total / 12
    from_phrases = max((int(x) for x in YEARS_PHRASE.findall(text)), default=0)
    return round(min(40.0, max(from_ranges, from_phrases)), 1)


def parse_resume(text: str) -> dict:
    tax = get_taxonomy()
    sections = split_sections(text)
    found_sections = [s for s in SECTION_KEYWORDS if s in sections]
    # Experience: prefer the experience section; never count education date ranges.
    exp_text = sections.get("experience") or "\n".join(v for k, v in sections.items() if k not in ("education", "header"))
    education_text = sections.get("education", "")
    education = [ln.strip(" -•\t") for ln in (education_text or text).splitlines() if DEGREE_RX.search(ln)][:5]
    words = re.findall(r"\b\w+\b", text)
    return {
        "skills": tax.extract_skills(text),
        "sections": found_sections,
        "years_experience": estimate_years(exp_text),
        "education": education,
        "word_count": len(words),
        "_sections_text": sections,
    }


# ---------------------------------------------------------------- scoring
def analyse(text: str, parsed: dict) -> dict:
    sections = parsed["sections"]
    lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
    bullets = [ln for ln in lines if re.match(r"^[-•*▪◦●]\s*", ln)]
    lower = text.lower()
    breakdown, strengths, suggestions = [], [], []

    def add(criterion, score, max_score, detail):
        breakdown.append({"criterion": criterion, "score": round(score, 1), "max": max_score, "detail": detail})

    # 1. Contact information (10)
    has_email, has_phone, has_link = bool(EMAIL_RX.search(text)), bool(PHONE_RX.search(text)), bool(LINK_RX.search(text))
    add("Contact information", 4 * has_email + 3 * has_phone + 3 * has_link, 10,
        f"email {'✓' if has_email else '✗'}, phone {'✓' if has_phone else '✗'}, profile link {'✓' if has_link else '✗'}")
    if not has_email:
        suggestions.append("Add a professional email address at the top.")
    if not has_link:
        suggestions.append("Add a GitHub, LinkedIn or portfolio link so recruiters can see your work.")

    # 2. Section completeness (25)
    weights = {"skills": 7, "experience": 6, "projects": 5, "education": 5, "summary": 2}
    sec_score = sum(w for s, w in weights.items() if s in sections)
    if "experience" not in sections and "projects" in sections:
        sec_score += 3  # students: strong projects partly substitute for experience
    missing = [s for s in weights if s not in sections]
    add("Section completeness", min(25, sec_score), 25,
        "found: " + (", ".join(sections) or "none") + ("; missing: " + ", ".join(missing) if missing else ""))
    for s in missing:
        suggestions.append({
            "skills": "Add a dedicated SKILLS section listing your tools and technologies.",
            "experience": "Add an EXPERIENCE section (internships, freelance or part-time work count).",
            "projects": "Add a PROJECTS section with 2-3 projects, the stack used and your impact.",
            "education": "Add an EDUCATION section with your degree, institution and graduation year.",
            "summary": "Add a 2-3 line SUMMARY tailored to the role you want.",
        }[s])

    # 3. Skills breadth (20)
    n_skills = len(parsed["skills"])
    add("Skills detected", min(20, n_skills * 2), 20, f"{n_skills} recognised skills")
    if n_skills >= 8:
        strengths.append(f"Broad technical toolkit ({n_skills} recognised skills).")
    elif n_skills < 5:
        suggestions.append("List more concrete technologies (languages, frameworks, databases, tools).")

    # 4. Impact & quantification (20)
    quantified = [b for b in bullets if NUMBER_RX.search(b)]
    verbs = {w for w in re.findall(r"[a-z]+", lower) if w in ACTION_VERBS}
    impact = min(15, len(quantified) * 5) + (5 if len(verbs) >= 3 else 2 if verbs else 0)
    add("Impact & quantified results", impact, 20,
        f"{len(quantified)} bullet(s) with numbers, {len(verbs)} strong action verb(s)")
    if quantified:
        strengths.append("Uses numbers to show impact.")
    else:
        suggestions.append("Quantify results (e.g. 'reduced load time by 40%', 'served 2,000 users').")
    if len(verbs) < 3:
        suggestions.append("Start bullet points with action verbs such as Built, Led, Improved or Automated.")

    # 5. Length (10)
    wc = parsed["word_count"]
    length_score = 10 if 300 <= wc <= 900 else 6 if 150 <= wc < 300 or 900 < wc <= 1300 else 2
    add("Length", length_score, 10, f"{wc} words (ideal 300-900 for students and early-career)")
    if wc < 300:
        suggestions.append("Your resume is short — add detail about projects and responsibilities.")
    elif wc > 1300:
        suggestions.append("Your resume is long — keep it to one or two pages of the most relevant work.")

    # 6. Readability (15)
    first_person = len(re.findall(r"\b(i|me|my)\b", lower))
    long_lines = [ln for ln in lines if len(ln.split()) > 40]
    readability = (5 if len(bullets) >= 3 else 2 if bullets else 0) + (5 if not long_lines else 2) + (5 if first_person <= 3 else 2)
    add("Readability", readability, 15,
        f"{len(bullets)} bullet points, {len(long_lines)} very long line(s), {first_person} first-person pronoun(s)")
    if len(bullets) < 3:
        suggestions.append("Use bullet points for responsibilities and achievements — they are easier to scan.")
    if first_person > 3:
        suggestions.append("Avoid 'I/me/my' in bullet points; start with the action instead.")

    if "projects" in sections:
        strengths.append("Includes a projects section — great for showing practical skills.")
    if parsed["years_experience"] >= 1:
        strengths.append(f"About {parsed['years_experience']} year(s) of experience detected.")

    score = round(sum(b["score"] for b in breakdown))
    return {"score": max(0, min(100, score)), "breakdown": breakdown, "strengths": strengths, "suggestions": suggestions[:8]}


def analyse_file(data: bytes, filename: str) -> dict:
    text = extract_text(data, filename)
    if len(text) < 50:
        raise ResumeError("Could not find enough text. If this is a scanned image, upload a text-based PDF or DOCX.")
    parsed = parse_resume(text)
    analysis = analyse(text, parsed)
    parsed.pop("_sections_text", None)
    return {"text": text, **parsed, "analysis": analysis}
